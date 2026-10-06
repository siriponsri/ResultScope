"""Optional report builder. Requires python-docx and Pillow; run from any directory."""
from pathlib import Path
import json,html
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from PIL import Image,ImageDraw,ImageFont
R=Path(__file__).resolve().parents[1];D=R/'docs/business-v3';out=D/'ResultScope_Business_Report.docx'
# Precise engineering diagram, not a generated scientific image.
W,H=1440,810;im=Image.new('RGB',(W,H),'white');draw=ImageDraw.Draw(im)
font='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf';f=ImageFont.truetype(font,23);fb=ImageFont.truetype('/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf',25)
boxes={'web':(30,25,440,150),'line':(1000,25,1410,150),'api':(510,210,930,345),'db':(30,415,440,550),'agent':(510,415,930,550),'tools':(1000,415,1410,550),'rag':(510,640,930,775),'payment':(1000,640,1410,775)}
labels={'web':('Web customer and staff','English UI / multilingual chat'),'line':('LINE webhook and worker','Signed events / durable jobs'),'api':('FastAPI business API','Session / ownership / version'),'db':('Encrypted business storage','PostgreSQL or local SQLite'),'agent':('Guarded LLM orchestration','Input / planner / answer / review'),'tools':('Confirmed action gateway','Price / capacity / authorization'),'rag':('Public evidence retrieval','BM25 + optional LightRAG'),'payment':('External sandbox services','Stripe / LINE / map provider')}
def arrow(a,b):
 draw.line([a,b],fill='#716b79',width=4)
 import math
 angle=math.atan2(b[1]-a[1],b[0]-a[0]);p=[b,(b[0]-15*math.cos(angle-.4),b[1]-15*math.sin(angle-.4)),(b[0]-15*math.cos(angle+.4),b[1]-15*math.sin(angle+.4))];draw.polygon(p,fill='#716b79')
for a,b in [((440,90),(560,210)),((1000,90),(880,210)),((565,345),(310,415)),((720,345),(720,415)),((880,345),(1130,415)),((720,550),(720,640)),((1205,550),(1205,640))]:arrow(a,b)
for k,(x1,y1,x2,y2) in boxes.items():
 draw.rounded_rectangle((x1,y1,x2,y2),radius=16,fill='#f8f8fc',outline='#c2bbc8',width=2)
 title,sub=labels[k]
 for j,(text,ff) in enumerate([(title,fb),(sub,f)]):
  w=draw.textlength(text,font=ff);draw.text(((x1+x2-w)/2,y1+29+j*42),text,font=ff,fill='#21172f')
diag=D/'architecture.png';im.save(diag)
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.75)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
 st=doc.styles[name];st.font.name='DejaVu Sans';st.font.color.rgb=RGBColor(0,0,0);st.font.size=Pt(10.5 if name=='Normal' else 11)
 st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.12
for name,size in [('Title',28),('Subtitle',14),('Heading 1',20),('Heading 2',13),('Heading 3',11)]:doc.styles[name].font.size=Pt(size)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(13);doc.styles['Heading 2'].paragraph_format.space_before=Pt(10)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=footer.add_run('ResultScope 3  |  ');r.font.size=Pt(8)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
doc.core_properties.title='ResultScope Business System Report';doc.core_properties.author='ResultScope project';doc.core_properties.subject='Intelligent Chatbot Development final project'
def p(text,style=None):return doc.add_paragraph(text,style)
def numbered(items):
 for i,text in enumerate(items,1):
  pp=p(str(i)+'. '+text);pp.paragraph_format.left_indent=Inches(.22);pp.paragraph_format.first_line_indent=Inches(-.22)

def h(text):return doc.add_heading(text,1)
def h2(text):return doc.add_heading(text,2)
def page(title):doc.add_page_break();h(title)
def table(headers,rows,widths=None):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 if widths:
  for col,w in zip(t.columns,widths):col.width=Inches(w)
 for i,v in enumerate(headers):t.rows[0].cells[i].text=str(v)
 for row in rows:
  c=t.add_row().cells
  for i,v in enumerate(row):c[i].text=str(v)
 for ri,row in enumerate(t.rows):
  pr=row._tr.get_or_add_trPr();cs=OxmlElement('w:cantSplit');pr.append(cs)
  if ri==0:pr.append(OxmlElement('w:tblHeader'))
  for ci,c in enumerate(row.cells):
   if widths:c.width=Inches(widths[ci])
   c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cp=c._tc.get_or_add_tcPr()
   borders=OxmlElement('w:tcBorders')
   for e in ['top','left','bottom','right']:
    b=OxmlElement('w:'+e);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
   cp.append(borders);m=OxmlElement('w:tcMar')
   for e in ['top','left','bottom','right']:
    b=OxmlElement('w:'+e);b.set(qn('w:w'),'70');b.set(qn('w:type'),'dxa');m.append(b)
   cp.append(m);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'263747' if ri==0 else 'F3F5F7' if ri%2==0 else 'FFFFFF');cp.append(sh)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(1);pp.paragraph_format.line_spacing=1.05
    for rr in pp.runs:rr.font.size=Pt(9);rr.font.bold=ri==0;rr.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
 p('')
 return t
# 1
p('ResultScope Business System Report','Title');p('Intelligent Chatbot Development 06048308','Subtitle');p('Version 3.0.0  |  5 October 2026  |  Coursework implementation')
p('ResultScope brings package selection, laboratory-report questions, appointments and staff support into one conversation. It serves individuals and organizations through a website and a LINE adapter. The LLM plans the conversation; software protects identity, clinical field fidelity and business transactions.')
p('The source package implements the full coursework business flow with realistic own-price packages and mock branches. Medical evidence is drawn from real public references. A model answer never establishes a diagnosis, and a simulated business record never becomes clinical reference knowledge.')
table(['Delivery area','Included'],[['Application','Customer workspace, staff desk, catalog, booking, sandbox payments, report history and LINE worker'],['Evidence','58 runtime records; 68-URL discovery manifest; original source files; six synthetic test reports'],['Engineering verification','Offline Python/API checks and browser flows; exact results in the evidence folder'],['Not established','Live provider quality, external deployment, clinical approval or completed course submission']],[1.6,5.4])
h2('Reading guide')
table(['Pages','Subject'],[['2–3','Business model, service catalog and customer experience'],['4–6','Architecture, one-message flow and medical evidence'],['7–9','Guardrails, channel operations, models and deployment'],['10–12','UAT, submission handoff and sources']],[.8,6.2])
p('Prepared for the project owner and local test agent. Student names, IDs and actual individual contributions must be completed by the project team before submission. AI-assisted implementation is disclosed separately from student authorship.')
#2
page('Business model and service catalog')
p('The center earns revenue from individual health checks, appropriate follow-up services and organization contracts. The product aims to reduce repeated questions and make package inclusions, evidence and next steps inspectable. These are product goals, not measured commercial outcomes.')
p('HealthLab informs service and branch clarity, BRIA informs package comparison and booking, and N Health informs specimen and turnaround metadata. ResultScope combines these ideas in an original conversation flow; it does not claim affiliation or proven superiority.')
cat=json.loads((R/'business_data/catalog.json').read_text())['packages']
table(['ID','Package or service','THB','Unit'],[[x['id'],x['name'],f"{x['price_thb']:,}",x['price_unit']] for x in cat],[.45,3.6,.7,1.25])
p('The 18 entries satisfy the assignment service-count route. Full inclusions and staff-review flags are in business_data/catalog.json. Prices are simulated business data, not competitor quotations. Follow-up services flagged for review cannot be booked automatically. Organization quotations require staff coordination.')
#3
page('Customer experience and business policy')
h2('Individual and returning customer')
p('A customer describes goals and budget in their own words. The assistant asks for missing context, compares current inclusions and proposes a quote. Booking requires a branch, date and time plus explicit confirmation. A returning customer can select confirmed current and previous reports; the assistant must not silently combine people, units or methods.')
h2('Organization and staff journey')
p('An organization contact supplies headcount, venue, preferred dates and service needs. Staff take over the conversation and prepare a quotation. The customer reviews and accepts it; payment remains a separate state. Onsite feasibility is coordinated by staff. There is no HR portal or employer entitlement to employees’ medical reports.')
table(['FAQ topic','Required response behavior'],[['Package fit','Ask about goals, budget and existing inclusions.'],['Prices','Use the current catalog and server-calculated total.'],['Branch and time','Use branch data and valid available slots.'],['Preparation','Use verified instructions; never advise medication changes.'],['Booking changes','Confirm the request and apply the stated policy.'],['Payment','Use test checkout or an auditable center receipt.'],['Result timing','Give a supported service estimate or refer to staff.'],['Flagged results','Explain confirmed fields with real medical sources.'],['Additional tests','Explain reason, uncertainty, overlap and optionality.'],['Organization visits','Collect logistics and route to a coordinator.']],[1.4,5.6])
p('The website is English; the chatbot is prompted to follow the user’s language. Multilingual effectiveness is an acceptance target that still needs live measurement. The purple visual identity, self-hosted fonts, source links and six supplied laboratory photographs support a calm, professional experience. Motion is finite and respects reduced-motion settings.')
#4
page('As built system architecture')
p('FastAPI, Jinja and vanilla JavaScript are retained. / is the business homepage, /app the customer workspace, /staff the service desk and /lab the retained laboratory workspace. The model is not a keyword router; it produces a constrained plan and a cited answer.')
doc.add_picture(str(diag),width=Inches(7));p('Figure 1. Implemented component boundaries. External connections require deployment credentials and live acceptance.',style='Caption')
table(['Component','Concrete responsibility'],[['business_agent.py','Plan intent, select retrieval, draft an answer, review support and apply guards.'],['business.py','Authorize, manage context and validate all explicit business actions.'],['business_store.py','Encrypted entity payloads, canonical prices, ownership and atomic transactions.'],['business_worker.py','Persisted LINE jobs, lease handling, identity rechecks and outbox delivery.'],['evidence_search.py','Local lexical ranking and optional known-ID LightRAG rank fusion.']],[1.6,5.4])
p('Cloud deployments require PostgreSQL. Local development uses SQLite. A database mutex serializes small-pilot state transitions; provider calls run outside it. This favors correctness over large-scale throughput. Original documents are encrypted in the database, with 3 MB upload and 20-report limits; object storage and load optimization remain future work.')
#5
page('One message and retrieval flow')
numbered(['The API validates the session and reserves a conversation version. Staff mode saves a message to the staff conversation without calling the bot.','The input guard checks the request. A structured LLM planner chooses intent, language, a medical search query and any proposed business action.','Business facts come from the canonical catalog. Medical evidence comes from local BM25 and, when configured, LightRAG mix graph/vector rankings.','The LLM drafts the answer. Schema, citation IDs and exact confirmed report fields are checked in code. A separate model call reviews support.','The output guard screens the final answer and proposed action. Only checked text is released. Stop, staff takeover or a version change discards stale work.','A transaction preview is not execution. On confirmation, the API rechecks ownership, expiry, current price, capacity and idempotency, then records the outcome.'])
h2('How the knowledge collection is built')
p('The supplied Deep Research document is retained as a discovery input. Its 68 explicit URLs were attempted; 64 HTTP bodies were downloaded, including PDFs, HTML pages and non-document responses. The manifest records URL, final URL, retrieval time, bytes, SHA-256 and status. A successful HTTP request is not automatically valid medical evidence.')
p('The runtime uses only knowledge/evidence/catalog.json: 35 earlier verified records plus 23 new source-checked intervals. The LightRAG export contains these same 58 records and stable resultscope://evidence/ IDs. Remote results contribute ranking only; arbitrary remote text cannot become a cited source.')
p('Customer documents, synthetic samples, expected_results.json, unresolved research citations and unreviewed raw guidelines stay outside shared RAG. An embedding-model or dimension change requires an index rebuild and fresh retrieval evaluation. Live semantic quality and ingestion cost are not established by offline source checks.')
#6
page('Medical evidence and exact values')
p('Reference intervals depend on the laboratory, specimen, assay, population and age. They are not interchangeable with diagnostic decision limits or treatment targets. A confirmed patient report’s interval takes precedence over a public manual. Missing or ambiguous intervals remain unknown; source checking does not constitute clinical approval.')
rows=json.loads((R/'knowledge/medical_sources/verified_intervals.json').read_text())['records']
table(['Test','Source population','Interval','Source'],[[x['test'],x['population']+'; '+x['sex_as_printed'],x['lower']+'–'+x['upper']+' '+x['unit'],x['source_id']+' p.'+str(x['source_page'])] for x in rows if x['test'] in ['WBC','HGB','Calcium','Magnesium','TSH','FT4','Ferritin','Vitamin B12']],[1.05,2.4,2.4,1.15])
p('These examples are from Siriraj Hospital primary documents. The hematology table is explicitly dated 29 March 2016; the other selected manuals are undated and retrieved on 5 October 2026. Full method/specimen notes, source links, hashes and all 23 records are in verified_intervals.json. They are contextual educational references, not automatic patient thresholds.')
h2('Corrections from source inspection')
p('The Iowa handbook at the discovery link is internally dated 2013, so it is excluded as a current numeric source. The WHO download returned an HTML shell, not its guideline PDF. The Siriraj immunology index was a maintenance page. CLSI provided a public landing page, not the paid standard. Four URLs failed or exceeded the download cap. No values were guessed to fill these gaps.')
p('OCR uses actual image/PDF content and requires field confirmation. The six supplied reports are explicitly synthetic test artifacts, not real patient records or clinical evidence. Report labels and the evaluator oracle retain that distinction.')
#7
page('Guardrails and Week 11 acceptance')
p('The architecture follows the instructor’s separate input and output guard approach. Llama Guard classifies configured hazards; it does not establish factual correctness or solve prompt injection alone. Exact-value checks, source identity, ownership, safe rendering and tool authorization provide separate controls.')
table(['Case','Required guard-service observation'],[['G01','GET /health returns 200 and status ok.'],['G02','Safe input returns exact message, input direction and empty categories.'],['G03','Unsafe input returns unsafe and the configured category.'],['G04','Output direction is preserved and classified.'],['G05','Configured blocklist canary is blocked without an upstream call.'],['G06','Missing/empty message or invalid direction returns 422.'],['G07','Isolated invalid upstream key returns 502 guard_unavailable; never safe.'],['G08','Allowed CORS preflight succeeds; disallowed origin receives no permission.']],[.7,6.3])
p('The supplied Postman collection targets a separately configured teacher-compatible /check service. It is not a fictitious business API endpoint. G05 needs an isolated canary configuration; G07 needs a separate bad-key instance. CORS must also be tested in a browser. These live cases remain NOT_RUN until actual evidence is recorded.')
h2('Threats and limits')
p('The Week 11 checklist maps prompt injection, sensitive disclosure, excessive agency, supply-chain risk, poisoning, unbounded consumption, misinformation, hidden-context exposure, vector weaknesses and improper output handling to concrete controls. Retrieved documents are untrusted data, not instructions. The bot cannot grant itself staff privileges, settle an order from a slip image or read another account.')
p('Guard outages, malformed verdicts and unsupported drafts fail closed. No unguarded model tokens are streamed. A same-model critic can share the generator’s errors; a fluent response with a source chip is not enough. Include benign Thai price questions to measure false refusals as well as adversarial prompts.')
#8
page('Transactions channels and privacy')
h2('Bookings and payment')
p('Half-hour individual slots use Asia/Bangkok, Monday–Saturday, within 30 days. Capacity, canonical prices and idempotency are checked inside a database transaction. Unconfirmed previews expire and are invalidated by price or context changes. Confirmed orders retain agreed totals. Follow-up services marked for review and all organization contracts require staff.')
p('Stripe accepts test-mode keys and signed events only. A matching checkout ID, amount, currency and booking are required before payment is marked paid. A redirect or uploaded receipt cannot authorize settlement. Managers approve refunds; provider refunds use a stable idempotency key. Pay-at-center receipt/refund records are clearly simulations.')
h2('LINE and staff')
p('LINE signatures are verified over raw webhook bytes. Events are deduplicated and persisted before acknowledgment. A worker processes text/images and shares the same business orchestration. Expired reply tokens require separately enabled push delivery. Uncertain model or send outcomes are not blindly replayed.')
p('Account linking requires a one-use short-lived invitation, recent website sign-in and explicit consent. Display names are not identity proof. Linking imports the verified LINE guest’s records; unlink cancels pending delivery and the worker rechecks ownership before sending. A network send already in flight cannot be recalled.')
p('Staff take over before replying; the bot pauses and stale in-flight drafts are discarded. Foreground polling updates new messages without replacing the staff reply draft. Managers can edit package price/availability and inspect delivery failures. Clinical is currently a staff-class role, not a clinical sign-off workflow.')
h2('Private data')
p('Customer payloads use application-level Fernet encryption; session tokens are random and stored by digest. Passwords use salted PBKDF2. Database metadata remains queryable. Report deletion clears current and archived conversational copies. Automated retention, account erasure, email verification/recovery and fine-grained clinical segregation require further production work. No regulatory or clinical readiness claim is made.')
#9
page('Models hosting and cost')
table(['Role','Selected approach','Important limit'],[['Conversation','Typhoon v2.5 30B A3B default','Use a bounded authorized cycle; multilingual quality needs measurement.'],['OpenThaiLLM','NECTEC 7B research candidate','Base model, not a drop-in hosted chatbot.'],['Alternative','iApp OpenThai 2.0 compatible adapter','Separate model family; verify current tariff and access.'],['OCR','Typhoon default; iApp optional','Evaluate exact fields on the same six images.'],['Guard','Llama Guard 4 or teacher service','Independent safety boundary; never disabled to save money.'],['Retrieval','BM25 plus optional LightRAG and BGE-M3','Pin embedding identity/dimension and approved source IDs.']],[1,2.8,3.2])
p('The preferred 0–300 THB plan is a limited coursework demo: eligible free web hosting, free database quotas, Typhoon access if available and a small guard/embedding/OCR allowance. It is not a promise of commercial uptime or unlimited free inference. iApp’s low per-page OCR tariff does not prove its minimum prepaid purchase fits this budget.')
p('A more stable course pilot can start from an estimated 1,000–1,500 THB/month, including a $25 compute sizing experiment and metered calls. This uses an assumed 35 THB/USD planning rate, not a current FX quote. Domain, taxes, paid LINE upgrades, clinical staffing and commercial high availability are excluded. No spending was authorized or incurred for model evaluation in this delivery.')
h2('Deployment choices')
p('Vercel runs the stateless Python API with external PostgreSQL and Redis. LightRAG and the LINE worker remain separate. Vercel Hobby is restricted to eligible non-commercial use. Render can run API plus worker, but its free service sleep and ephemeral storage affect reliability. Do not use an expiring free database for durable records.')
p('Preserve BUSINESS_DATA_KEY with its database, configure owner credentials and run the deployment acceptance checklist. The current budget control counts upstream attempts per named cycle; it is not a per-user THB billing engine. Count guard, planning, answer, critic, OCR and retrieval work separately.')
#10
page('Verification and acceptance evidence')
p('The delivered evidence folder separates automated engineering observations from real-provider outcomes. Python tests exercise business authorization, canonical prices, capacity races, idempotency, payment signatures, report ownership, account linking, guard failures and source integrity. Browser UAT uses real UI/API/storage with explicit LLM/OCR doubles.')
table(['Layer','Evidence','Meaning'],[['Python/API','evidence/pytest.txt and pytest.xml','Executed offline; use the recorded final count.'],['Browser','evidence/browser-uat.json and PNGs','Customer/staff flows, desktop/mobile and JavaScript checks.'],['Source integrity','evidence/source-integrity.json','Packaged bytes match recorded hashes; not clinical review.'],['Provider and cloud','tests/uat_cases.json','Live cases NOT_RUN until credentials and bounded scope are active.']],[1.1,2.7,3.2])
h2('Required live coursework cases')
table(['Set','Coverage','Evidence to record'],[['T01–T10','Budget, comparisons, branch, preparation, reports, booking, organization and language','Exact answer, citations, verdict and total elapsed time.'],['I01–I06','All six supplied report images','OCR fields, missing/extra rows, critical errors, latency and cost.'],['S01–S05','Agency, privacy, injection, clinical boundaries and document instructions','Actual guarded output, expected behavior and final verdict.'],['Extended cases','73-case total across business, retrieval, security and operations','PASS / FAIL / BLOCKED / NOT_RUN with artifact paths.']],[1,2.8,3.2])
p('An OCR double passing a review dialog does not measure OCR accuracy. An offline guard failure test does not measure multilingual safety. A cloud liveness check does not establish persistence or signed-event delivery. The six evaluator oracles are used only after extraction for scoring, never as model input.')
p('Critical failures include cross-customer disclosure, an unauthorized transaction, wrong critical numeric/unit/range/flag fields, unsafe unguarded output or credential exposure. Do not average these away with successful ordinary questions.')
#11
page('Local agent and submission handoff')
p('The local agent starts from the complete replacement ZIP, preserves existing private state and records the source identity. Windows verification uses scripts/check.ps1; this delivery’s execution environment is Linux, so a Windows pass must not be inferred.')
numbered(['Install dependencies and rerun compileall, pytest and the medical-source verifier. Run npm run uat:business with an isolated database; inspect screenshots, not just HTTP status.','Configure the intended model/guard/OCR and external services with an explicit attempt budget. Stop on an unavailable guard, invalid key or exhausted cycle; never reset counters to hide a failed attempt.','Run the ten text, six image, five safety and Week 11 cases. Record actual outputs, latencies, model IDs, source IDs and costs. Test PostgreSQL persistence, LINE signatures/linking, Stripe sandbox and maps separately.','Repeat three fixed cases before and after a meaningful prompt/retrieval/guard change. Keep model configuration and inputs comparable; preserve failures and regressions as evidence.','Complete student names, IDs, role contributions and references. Import the DOCX into Google Docs, verify its PDF, and record a demonstration no longer than 180 seconds.'])
h2('What remains for course submission')
p('The source package and report are prepared locally. Native Google Docs import, Drive folder submission and the final video are not claimed as completed. Actual live answers and matched before/after model improvements remain blank until measured. The templates intentionally do not invent successful outputs or student contributions.')
table(['Course date in Bangkok','Milestone from supplied material'],[['3 October 2026','Business-name notification date; no submission claim.'],['10 October 2026','Progress and examination.'],['17 October 2026','Submission.'],['24 October 2026','Presentation.']],[2,5])
p('The business simulation permission comes from the owner’s clarification. It does not permit invented medical sources. The Final Project, Weeks 1–12 and Week 11 guard requirements are mapped in 07_COURSE_TRACEABILITY.md, with the original course files and hashes included.')
#12
page('Sources and reproducibility')
p('The complete register is docs/business-v3/08_SOURCE_REGISTER.md. Medical URLs, document pages, retrieval times and hashes are in knowledge/medical_sources/source_manifest.json and verified_intervals.json. Original course files are listed in course-source-manifest.json. Access observations were made on 5 October 2026; deployment tariffs should be rechecked before activation.')
refs=[('Business references','healthlabclinic.com/en/; booking.brianet.com/product-category/all-package/; weblabonline.nhealth-asia.com/website/Catalog'),('Medical authority','Faculty of Medicine Siriraj Hospital, Mahidol University laboratory manuals; MedlinePlus, U.S. National Library of Medicine. Exact URLs are preserved in the runtime catalogue.'),('Instructor repositories','github.com/chacharin/light-rag; render-rag; mcp-lightrag; llama-guard-layer'),('Retrieval implementation','github.com/HKUDS/LightRAG. Reviewed README blob 5dd8b7db99231f6ecede9b76ba8dc304d9638f3c; PGTableGraphStorage blob 179cea181dc31e078b0f016267928908d05ee4ba.'),('Models and OCR','docs.opentyphoon.ai; huggingface.co/nectec/OpenThaiLLM-Prebuilt-7B; iapp.co.th/docs/ocr/document; huggingface.co/meta-llama/Llama-Guard-4-12B; huggingface.co/BAAI/bge-m3.'),('Platform boundaries','render.com/docs/free; vercel.com/docs/plans/hobby; supabase.com/pricing; developers.line.biz/en/docs/messaging-api/; docs.stripe.com/webhooks.'),('Visual references','Bloom Dribbble shot 26661884 and Genomic shot 24697057 by Phenomenon Studio. Design principles only; no copied Dribbble artwork.'),('2035 research scenario','ReAct, openCHA, AMIE, ALCE, Self-RAG, LightRAG and Generative Agents. Full primary-paper citations and limitations are in docs/research/CHATBOT_2035.md.')]
for name,text in refs:
 hh=h2(name);hh.paragraph_format.space_before=Pt(6);hh.paragraph_format.space_after=Pt(4)
 pp=p(text);pp.paragraph_format.space_after=Pt(5)
p('The 2035 direction is an engineering scenario: inspectable memory, bounded tool use, evidence-linked conversation and smooth human takeover. It is not a prediction, a reproduction of the cited research systems or evidence that autonomous clinical care is safe.')
p('Reproduce with scripts/package_v3.py. DELIVERY_MANIFEST.json records source hashes. Preserve licenses, image provenance, evidence separation and the usage ledger.')
doc.save(out);print(str(out))
